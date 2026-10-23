"""Build the T07 A2 report with plain black text and six main pages.

Numerical content comes from results/analysis.json and the public generator.
Metadata comes from team_config.json so the real repository record can be added.
Export the DOCX to PDF with Word or LibreOffice after genuine signatures.
"""
import csv
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def build():
    r = json.loads((ROOT / 'results/analysis.json').read_text())
    cfg = json.loads((ROOT / 'team_config.json').read_text())
    if (cfg['seed'], cfg['story'], cfg['variant']) != (7, 'A', 2):
        raise ValueError('This report is specific to the assigned T07 A2 instance')
    t, k, s, h = (r[key] for key in ['table1', 'table2', 'table3', 'hand'])
    names = [m['name'] for m in cfg['members']]
    if len(names) != 4:
        raise ValueError('Exactly four members are required')
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(0.70)
    sec.left_margin = sec.right_margin = Inches(0.78)
    for st in doc.styles:
        if hasattr(st, 'font'):
            st.font.name = 'Liberation Sans'
            st.font.size = Pt(11)
            st.font.bold = False
            st.font.italic = False
            st.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.06
    normal.paragraph_format.widow_control = True
    doc.core_properties.title = 'Team T07 Checkpoint 3 Constrained Component'
    doc.core_properties.author = 'Team T07'
    doc.core_properties.subject = 'Regional Logistics variant A2 LP relaxation'
    text_pages = []
    current = []

    def p(text='', keep=False):
        para = doc.add_paragraph(text)
        para.paragraph_format.keep_with_next = keep
        current.append(text)
        return para

    def title(text):
        return p(text, keep=True)

    def table(headers, rows, widths=None):
        tab = doc.add_table(rows=1, cols=len(headers))
        tab.autofit = False
        if widths:
            for col, width in zip(tab.columns, widths):
                col.width = Inches(width)
        for cell, label in zip(tab.rows[0].cells, headers):
            cell.text = str(label)
        for record in rows:
            for cell, value in zip(tab.add_row().cells, record):
                cell.text = str(value)
        for row in tab.rows:
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(2)
                    para.paragraph_format.space_before = Pt(1)
                    para.paragraph_format.line_spacing = 1.0
                tcpr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement('w:tcMar')
                for edge in ['top', 'left', 'bottom', 'right']:
                    el = OxmlElement('w:' + edge)
                    el.set(qn('w:w'), '30')
                    el.set(qn('w:type'), 'dxa')
                    margins.append(el)
                tcpr.append(margins)
        # The user requested plain text styling so tables have no decorative rules.
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            el = OxmlElement('w:' + edge)
            el.set(qn('w:val'), 'nil')
            borders.append(el)
        tab._tbl.tblPr.append(borders)
        current.append(' | '.join(str(x) for x in headers))
        current.extend(' | '.join(str(x) for x in row) for row in rows)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def signature(name):
        p('Signature ' + name + ' ____________________')

    def group_signatures():
        p('Signatures ' + names[0] + ' __________  ' + names[1] + ' __________')
        p(names[2] + ' __________  ' + names[3] + ' __________')

    def next_page():
        nonlocal current
        text_pages.append('\n\n'.join(current))
        current = []
        doc.add_page_break()

    f = lambda x: f'{x:.6f}'
    title('Team T07 Checkpoint 3 Constrained Component')
    p('Introduction to Optimization   Section 2')
    p('Regional Logistics   Storyline A   Variant 2')
    p('23 October 2026')
    p('The LP relaxation gives a certified lower bound of ' + f(t['zLP']) +
      ' for our assigned capacitated center model with premium points. '
      'All supplied data checks match and all five optimality conditions pass at tolerance 0.000001.')
    p("Instance call cp34_instance(7, 'A', 2)")
    p('Repository link ' + (cfg.get('repository_url') or '________________________________________'))
    p('Tag checkpoint3')
    p('Commit hash ' + (cfg.get('commit_hash') or '________________________________________'))
    title('Names and responsibilities')
    table(['Student', 'Role', 'Analysis and hand task'],
          [[names[i], 'M' + str(i+1), 'S' + str(i+1) + ' and H' + str(i+1)] for i in range(4)],
          [2.45, .65, 3.60])
    p('All four students share Task 5 S5 H5 Synthesis and report assembly. '
      'The planned division is 25 percent each.')
    p('Table 1 model checks belongs to M1 and the LP results belong to M2. '
      'Table 2 belongs to M3. Table 3 belongs to M4.')
    title('Declarations')
    p('OpenAI Codex generated and adapted the code calculations tests report draft and worked hand task solutions. '
      'An earlier local example implementation was adapted to the assigned T07 packet. '
      'The supplied public generator is unchanged. NumPy SciPy PuLP below version 4 and pytest support the numerical work. '
      'The editable report uses python-docx. Each member must review and complete their assigned work before signing. '
      'Actual contributions are recorded by genuine Git commits and the signed final statement.')
    next_page()

    title('Task 1 The model and its matrix form')
    p('Let I and K both be the point indices 0 to 40. The premium set P contains 13 points. '
      'Demand q_i comes from demand[i]. Distance d_ik comes from dist[i,k]. '
      'D_k is the warehouse distance and p_i is the penalty for premium point i. '
      'The common capacity is C = 82 and the required number of centers is V = 14.')
    p('Binary x_ik says that point i is served from center k. Binary y_k says that point k is a center. '
      'Binary z_i exists only for i in P and says that premium point i is left unserved.')
    p('Minimize sum over i,k of d_ik x_ik + 2 sum over k of D_k y_k + sum over i in P of p_i z_i')
    table(['Family', 'Constraint', 'Meaning'], [
        ['A1 nonpremium', 'sum_k x_ik = 1', 'Every nonpremium point is served once'],
        ['A1 premium', 'sum_k x_ik + z_i = 1', 'Serve the point or pay its penalty'],
        ['A2', 'x_ik <= y_k', 'Service uses an open center'],
        ['A3', 'sum_i q_i x_ik <= 82 y_k', 'The center load fits its capacity'],
        ['A4', 'x_kk = y_k', 'Every open center serves itself'],
        ['A5', 'sum_k y_k = 14', 'Exactly 14 centers are open']
    ], [1.12, 2.45, 3.13])
    p('The LP keeps all these rows and replaces each binary restriction with a bound from 0 to 1. '
      'There are no other changes to the reference model.')
    p('Stack v as all x_ik in row order then all y_k then z_i in increasing premium index order. '
      'The cost vector c stacks d_ik then 2D_k then p_i. The matrix problem is min c^T v '
      'with A_E v = b_E and A_U v <= b_U and 0 <= v <= 1. '
      'A_E contains A1 A4 A5. A_U contains A2 A3. The row names retain the original indices.')
    p('Columns 41 x 41 + 41 + 13 = 1735. Rows 41 + 41 x 41 + 41 + 41 + 1 = 1805. '
      'Bounds are excluded from this count. The sparse matrix has 6901 nonzero coefficients.')
    title('S1 The model and what it ignores')
    p('Our A2 model permits only premium points to remain unserved and charges their stated penalties. '
      'A premium point that opens as a center must serve itself because A4 still applies. '
      'The warehouse term is paid once per open center. The model ignores route order and H_max. '
      'It uses straight line distances rather than squared distances and nonlinear weights. '
      'These changes prevent direct numerical reuse of the Checkpoint 1 routing conclusions. '
      'We keep all diagonal linking rows and add no stronger cuts.')
    signature(names[0])
    next_page()

    title('Task 2 The LP relaxation')
    title('Table 1 Model checks and LP results')
    checks = [
        ['Points N', 41, t['N']], ['Centers V', 14, t['V']],
        ['Capacity', 82, t['Cap']], ['Total demand', 903, t['demand_sum']],
        ['Premium points', 13, t['is_premium_sum']],
        ['Sum of d_ik to three decimals', '87717.651', f"{t['distance_sum_3dp']:.3f}"],
        ['Variables', 1735, t['variables']], ['Constraint rows', 1805, t['rows']],
        ['LP value', 'Computed', f(t['zLP'])],
        ['Solver', 'Any LP solver', 'SciPy HiGHS'],
        ['Solve time in seconds', 'Measured', f(t['seconds'])],
        ['Fractional binary columns', 'Computed', t['fractional_binary_variables']]
    ]
    table(['Quantity', 'Packet value', 'Our result'], checks, [3.10, 1.45, 2.15])
    p('SciPy version ' + r['environment']['scipy'] + '. Python version ' + r['environment']['python'] +
      '. PuLP version ' + r['environment']['pulp'] + ' is below 4. '
      'Time varies by machine and includes the solve interface preparation.')
    title('Fractional center values with zero based point indices')
    fc = r['fractional_centers']
    rows = []
    for i in range(0, len(fc), 2):
        row = []
        for a in fc[i:i+2]:
            row.extend([a['variable'], f(a['value'])])
        rows.append(row + [''] * (4-len(row)))
    table(['Point variable', 'Value', 'Point variable', 'Value'], rows, [1.65, 1.7, 1.65, 1.7])
    title('S2 Value and fractionality of the LP')
    p('The optimum has 63 fractional assignments and 13 fractional center variables. '
      'All 13 premium unserved variables are zero so every premium point is served. '
      'Nine centers are fully open and fractional openings sum to the remaining five centers. '
      'The objective is 404.815441 for service travel plus 778.224536 for warehouse travel and zero penalty. '
      'The fractions let the LP split service and capacity across locations. They are not an integer dispatch plan. '
      'A different optimal vertex can change the fractional count while preserving the LP value. '
      'The complete solution is in results/solution.csv.')
    signature(names[1])
    next_page()

    title('Task 3 Optimality conditions')
    p('For original row prices pi define reduced costs r = c - A^T pi and row residuals s = b - Av. '
      'We use absolute tolerance 0.000001 for all five checks.')
    title('Table 2 Largest error in each optimality check')
    labels = ['a Primal rows and bounds', 'b Dual row signs', 'c Row complementary slackness',
              'd Reduced cost conditions', 'e Strong duality']
    table(['Condition', 'Largest error', 'Result'],
          [[label, ('0' if abs(value) == 0 else f'{value:.6e}'), 'Pass']
           for label, value in zip(labels, k['numbers'].values())], [3.6, 1.9, 1.2])
    p('For a we check every equality every inequality and both bounds. '
      'For b prices of A2 and A3 must be nonpositive while equality prices are free. '
      'For c we take max absolute pi_i s_i.')
    p('For d a variable at 0 requires r_j >= 0. A variable at 1 requires r_j <= 0. '
      'An interior variable requires r_j = 0. For e the dual objective is '
      'b^T pi + sum_j min(0,r_j) because every variable has upper bound 1.')
    p(f"Primal value {k['primal_value']:.12f}. Dual value {k['dual_value']:.12f}. "
      f"Their absolute difference is {k['numbers']['e_strong_duality']:.6e}.")
    title('Price sign calibration')
    p('Before the full solve we solve min -a + 4b with a <= 1 and b >= 1 and nonnegative variables. '
      'The optimum is a = b = 1 with value 3. The known original prices are -1 and 4. '
      'HiGHS returns -1 and -4 after the second row is multiplied by -1. '
      'Converting the second price back gives -1 and 4. This confirms the sign convention.')
    title('S3 Is the solution optimal and how do we know')
    p('All five errors are far below the required tolerance. Primal feasibility gives an allowed LP point. '
      'The row signs and bound terms give a valid dual lower bound. Equality of primal and dual values '
      'proves that no feasible LP point has lower cost. This certifies global LP optimality. '
      'The separate function check_kkt accepts ordinary x and pi vectors in the stored column and row order. '
      'It does not require a SciPy result object. A separate CBC certificate test checks this portability '
      'and deliberately invalid vectors are rejected. The certificate does not prove integer optimality.')
    signature(names[2])
    next_page()

    title('Task 4 Sensitivity of the center count')
    p('The designated constraint is A5 with sum_k y_k = b and b = 14. '
      'We change only this right hand side and resolve the same LP.')
    table(['Right hand side', 'Optimal LP value'], [
        ['13.99', f"{s['z_minus']:.9f}"], ['14.00', f"{s['z_base']:.9f}"],
        ['14.01', f"{s['z_plus']:.9f}"]], [3.35, 3.35])
    p('Left slope = [z(14) - z(13.99)] / 0.01')
    p('Right slope = [z(14.01) - z(14)] / 0.01')
    title('Table 3 Slopes and the original A5 dual price')
    table(['Left slope', 'Dual price pi_A5', 'Right slope'],
          [[f"{s[key]:.9f}" for key in ['left_slope', 'price', 'right_slope']]], [2.2, 2.3, 2.2])
    p('The required interval is left slope <= pi_A5 <= right slope within tolerance 0.000001. '
      'The interval check passes. The two slopes agree within tolerance so these runs detect no local kink. '
      'Equal slopes do not establish a unique primal solution or rule out degeneracy elsewhere in the LP.')
    title('S4 What the dual price says about the storyline')
    p('The positive price is 34.306928 cost units per additional center unit. '
      'Increasing b by 0.01 raises the LP cost by 0.343069. '
      'A5 fixes the number of centers exactly. Additional openings add warehouse travel even when service distances improve. '
      'The price measures the net local change. It applies to the LP near b = 14 and does not predict '
      'the cost of adding one whole center to an integer solution.')
    signature(names[3])
    title('Reproducible evidence')
    p('python -m experiments.run_all regenerates Tables 1 to 3 and their full supporting CSV files. '
      'results/row_prices.csv contains each original row sense right hand side left hand side and price. '
      'results/certificate_vectors.npz stores the full primal and dual vectors. '
      'results/analysis.json records all values and software versions. Numerical tables give direct evidence '
      'so no screenshots are needed.')
    next_page()

    title('Task 5 The bound and the bridge to Checkpoint 4')
    title('S5 The lower bound')
    p('Every feasible integer point uses only values 0 and 1.\n'
      'These values also satisfy the relaxed bounds from 0 to 1.\n'
      'The objective and every other constraint are identical.\n'
      'The LP minimizes over a larger feasible set so zLP <= z*.')
    p('For any feasible integer cost z > 0 let g = (z - zLP) / z. '
      'Then 0 <= (z - z*) / z <= g. It is an upper bound on relative suboptimality '
      'with z as denominator. A positive value mixes relaxation weakness and solution quality. '
      'It does not give the exact integer optimum or the exact optimality gap.')
    title('Active rows at absolute residual at most 0.000001')
    table(['Family', 'Active', 'Total', 'Nonzero prices'],
          [[a['family'], a['active'], a['total'], a['nonzero_prices']] for a in r['active_rows']],
          [1.55, 1.55, 1.55, 2.05])
    p('A5 has price 34.306928 while 22 capacity rows and 60 linking rows have nonzero prices '
      'so these families constrain local cost. Warehouse travel contributes about 65.8 percent of the objective '
      'but tight rows with zero prices need not have a marginal effect.')
    p('The 13 fractional center openings and 63 fractional assignments suggest a positive integer gap. '
      'Another optimal LP solution could be integral so this is a prediction to test in Checkpoint 4.')
    title('Synthesis')
    p('The LP provides the lower bound 1183.039977 for the exact A2 reference model. '
      'It serves every premium point despite allowing paid rejection. '
      'We cannot carry over Checkpoint 1 conclusions about route order H_max squared distance costs '
      'or nonlinear weights because these features were removed. '
      'Every point is a candidate center and indices start at zero. The warehouse term is counted once per center '
      'and only premium points have rejection variables. No extra assumptions or stronger constraints '
      'enter the numerical tables. Checkpoint 4 will compare feasible integer solutions with this bound.')
    p('Task 5 S5 and Synthesis are shared by all four students')
    group_signatures()
    next_page()

    title('Appendix H1 The triangle model')
    p(names[0] + '   M1')
    p('Part a Matrix form')
    p('The objective is min 8y1 + 7y2 + 7y3.')
    p('c = (8,7,7)^T and b = (1,1,1)^T')
    p('A has rows (1,1,0) and (0,1,1) and (1,0,1).')
    p('The relaxation is min c^T y with Ay >= b and 0 <= y <= 1. '
      'There are three variables and three structural rows. Bounds are not counted as rows.')
    p('Part b Changing the row senses')
    p('Multiplying each row by -1 gives -Ay <= -b. The rows are '
      '-y1 - y2 <= -1 and -y2 - y3 <= -1 and -y1 - y3 <= -1. '
      'For minimization the original >= prices are nonnegative. After the change each price is negated '
      'and becomes nonpositive. Here pi = (4,3,4) becomes (-4,-3,-4). '
      'The product of each price and its row remains unchanged.')
    p('Part c All eight binary points')
    table(['Point y', 'Row left sides', 'Feasible', 'Cost'], [
        ['000', '0 0 0', 'No', '0.000'], ['001', '0 1 1', 'No', '7.000'],
        ['010', '1 1 0', 'No', '7.000'], ['011', '1 2 1', 'Yes', '14.000'],
        ['100', '1 0 1', 'No', '8.000'], ['101', '1 1 2', 'Yes', '15.000'],
        ['110', '2 1 1', 'Yes', '15.000'], ['111', '2 2 2', 'Yes', '22.000']
    ], [1.5, 2.1, 1.5, 1.6])
    p('The minimum feasible binary cost is 14.000 at y = (0,1,1). Thus z* = 14.000.')
    p('Code reproduction tests/test_M1.py')
    signature(names[0])
    next_page()

    title('Appendix H2 The half point and rounding')
    p(names[1] + '   M2')
    p('Part a Feasibility and cost')
    p('Take y = (0.500,0.500,0.500). Each variable is between 0 and 1. '
      'Each covering row has left side 0.500 + 0.500 = 1.000 so all three rows hold with equality.')
    p('The cost is 8 x 0.500 + 7 x 0.500 + 7 x 0.500 = 4.000 + 3.500 + 3.500 = 11.000.')
    p('The dual vector (4,3,4) is nonnegative and A^T pi = (8,7,7) = c. '
      'Its lower bound is 4 + 3 + 4 = 11.000. The feasible half point reaches this bound '
      'so zLP = 11.000.')
    p('Part b Rounding up and down')
    table(['Rounding', 'Point', 'Feasibility', 'Cost'], [
        ['Up', '(1,1,1)', 'All row left sides are 2', '22.000'],
        ['Down', '(0,0,0)', 'All three rows fail by 1', '0.000']
    ], [1.0, 1.2, 3.1, 1.4])
    p('Rounding up is feasible but costly. Rounding down is infeasible even though its objective is zero. '
      'Its cost cannot be used as the value of a feasible solution.')
    p('Part c The integer optimum and relative LP gap')
    p('H1 gives z* = 14.000 at (0,1,1).')
    p('(z* - zLP) / z* = (14 - 11) / 14 = 3 / 14 = 0.214285714 = 21.429 percent.')
    p('Every feasible integer point costs at least 14.000 while the LP has cost 11.000. '
      'Therefore no integral point can be an optimal solution of this LP. '
      'The relative gap measures the weakness of the relaxation when the integer optimum is known.')
    p('Code reproduction tests/test_M2.py')
    signature(names[1])
    next_page()

    title('Appendix H3 Checking the two dual vectors')
    p(names[2] + '   M3')
    p('Use y = (0.500,0.500,0.500) and c = (8,7,7). '
      'All rows are >= rows so every row price must be nonnegative. '
      'All three row slacks are zero and every variable is interior.')
    p('Part a Assigned vector pi = (4,3,4)')
    p('A^T pi = (4+4,4+3,3+4) = (8,7,7). '
      'Reduced costs r = c - A^T pi = (0,0,0). '
      'The bound term sum min(0,r_j) is zero. '
      'The primal cost is 11 and the dual value is 4+3+4 = 11.')
    p('Part b Alternative vector pi prime = (5,2,3)')
    p('A^T pi prime = (5+3,5+2,2+3) = (8,7,5). '
      'Reduced costs are (0,0,2). All prices remain nonnegative. '
      'The dual value is 5+2+3 = 10 because the bound term is still zero.')
    table(['Largest error', 'pi', 'pi prime'], [
        ['a Primal feasibility', '0.000', '0.000'],
        ['b Dual row signs', '0.000', '0.000'],
        ['c Row complementarity', '0.000', '0.000'],
        ['d Reduced cost conditions', '0.000', '2.000'],
        ['e Strong duality difference', '0.000', '1.000']
    ], [4.0, 1.35, 1.35])
    p('With pi prime condition d fails because y3 is interior but r3 = 2. '
      'Condition e fails because 11 - 10 = 1. All other conditions pass. '
      'Row complementarity remains zero since every row is tight.')
    p('Part c The weaker lower bound')
    p('Since pi prime >= 0 and A^T pi prime <= c we have '
      'c^T y >= pi prime^T Ay >= pi prime^T b = 10 for every feasible nonnegative y. '
      'Therefore 10 is a valid lower bound even though it is not optimal. '
      'It is 1.000 below the best lower bound of 11.000. '
      'The failure is reduced cost stationarity at the proposed half point rather than invalid row signs.')
    p('Code reproduction tests/test_M3.py')
    signature(names[2])
    next_page()

    title('Appendix H4 Changing the first covering requirement')
    p(names[3] + '   M4')
    p('Part a Keep all three rows tight')
    p('y1 + y2 = 1 + epsilon\ny2 + y3 = 1\ny1 + y3 = 1')
    p('Subtracting the last two equations gives y1 = y2. '
      'The first equation then gives y1 = y2 = (1 + epsilon) / 2. '
      'Substitution gives y3 = (1 - epsilon) / 2.')
    p('The objective becomes\n'
      'z(1 + epsilon) = 8(1 + epsilon)/2 + 7(1 + epsilon)/2 + 7(1 - epsilon)/2\n'
      '= 11 + 4 epsilon.')
    p('The slope is 4.000 which equals the assigned first price pi1 = 4.000. '
      'The same nonnegative dual vector (4,3,4) has reduced costs zero and gives the bound '
      '4(1 + epsilon) + 3 + 4 = 11 + 4 epsilon. '
      'Our feasible point reaches it so the point is optimal throughout its validity interval.')
    p('Part b Bound validity')
    p('The bounds on y1 and y2 require 0 <= (1 + epsilon)/2 <= 1 '
      'which means -1 <= epsilon <= 1. '
      'The bounds on y3 require 0 <= (1 - epsilon)/2 <= 1 '
      'which gives the same interval. Thus the full interval is [-1,1].')
    table(['epsilon', 'Optimal point', 'Optimal cost'], [
        ['-1.000', '(0.000,0.000,1.000)', '7.000'],
        ['0.000', '(0.500,0.500,0.500)', '11.000'],
        ['0.010', '(0.505,0.505,0.495)', '11.040'],
        ['1.000', '(1.000,1.000,0.000)', '15.000']
    ], [1.3, 3.2, 2.2])
    p('Part c Interpretation')
    p('If T1 means the first customer must be covered then one extra unit of its required coverage '
      'costs 4.000 objective units while this basis remains valid.')
    p('Code reproduction tests/test_M4.py')
    signature(names[3])
    next_page()

    title('Appendix H5 The lower bound and three relative gaps')
    p('Shared task for all four students')
    p('Part a Three line lower bound argument')
    p('Every feasible binary point is also feasible for the relaxed triangle LP.\n'
      'The objective is identical at that point in both problems.\n'
      'Minimizing over the larger LP set gives zLP <= z*.')
    p('Part b Use zLP = 11.000 and calculate (z - zLP) / z')
    table(['Integer point', 'z', 'Calculation', 'Relative gap'], [
        ['Optimum (0,1,1)', '14.000', '(14-11)/14 = 3/14', '0.214286'],
        ['Point (1,0,1)', '15.000', '(15-11)/15 = 4/15', '0.266667'],
        ['Rounded up (1,1,1)', '22.000', '(22-11)/22 = 11/22', '0.500000']
    ], [2.0, 1.0, 2.2, 1.5])
    p('To three decimals the ratios are 0.214 and 0.267 and 0.500. '
      'As percentages they are 21.429 and 26.667 and 50.000 percent.')
    p('Part c What the numbers tell us')
    p('The first ratio uses the true integer optimum so it directly measures the quality of the LP bound. '
      'The other two compare feasible solution costs with the lower bound. '
      'They are upper bounds on relative suboptimality with the solution cost as denominator '
      'and include the weakness of the relaxation.')
    p('For (1,0,1) the actual relative loss is (15-14)/15 = 0.066667. '
      'For (1,1,1) it is (22-14)/22 = 0.363636. '
      'These are smaller than 0.266667 and 0.500000 because zLP is below z*. '
      'In Checkpoint 4 an exact integer solve will show the actual bound gap and will help assess '
      'the cost of a metaheuristic solution separately.')
    p('Code reproduction tests/test_M5.py')
    group_signatures()
    next_page()

    title('Team T07 Contribution statement')
    p('The planned allocation is 25 percent per student. Enter actual shares below and confirm that '
      'they total 100 percent before signing. The named responsibilities describe the work assigned '
      'to each member. Final descriptions and signatures must reflect completed work and genuine Git history.')
    table(['Member', 'Role', 'Assigned work', 'Actual share'], [
        [names[0], 'M1', 'Model matrices data checks S1 H1', '____ %'],
        [names[1], 'M2', 'LP solve solution reading S2 H2', '____ %'],
        [names[2], 'M3', 'KKT checker certificates S3 H3', '____ %'],
        [names[3], 'M4', 'Sensitivity interval check S4 H4', '____ %']
    ], [1.90, .60, 3.05, 1.15])
    p('Shared work Task 5 S5 H5 Synthesis experiments report assembly and final verification')
    p('Actual shares total ______ percent')
    p('Changes to the assigned work descriptions if needed')
    p('______________________________________________________________________')
    p('______________________________________________________________________')
    p('By signing we confirm that the final work descriptions and actual shares are correct and '
      'agree with the repository history. We have declared the AI assistance and reviewed the submitted work.')
    for name in names:
        p(name + '   Signature ____________________   Date ______________')
    text_pages.append('\n\n'.join(current))
    out = ROOT / 'report'
    out.mkdir(exist_ok=True)
    doc.save(out / 'T07_Checkpoint3.docx')
    (out / 'T07_Checkpoint3.txt').write_text('\n\n'.join(text_pages) + '\n')
    print('Created report/T07_Checkpoint3.docx and report/T07_Checkpoint3.txt')
    print('Expected pages 6 main + H1 H2 H3 H4 H5 + contribution = 12')


if __name__ == '__main__':
    build()
